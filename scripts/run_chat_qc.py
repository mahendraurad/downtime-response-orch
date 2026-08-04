"""Run the 50-case chat QC pack through the local FastAPI application."""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
CASES_PATH = ROOT / "quality" / "chat_qc_questions.json"
DEFAULT_OUTPUT = ROOT / "quality" / "results" / "DRO_Chat_QC.xlsx"


def _evaluate(case: dict, status_code: int, payload: dict) -> tuple[bool, list[str]]:
    failures: list[str] = []
    if status_code != 200:
        failures.append(f"HTTP {status_code}")
    if case.get("expected_intent") and payload.get("intent") != case["expected_intent"]:
        failures.append(
            f"intent {payload.get('intent')!r} != {case['expected_intent']!r}"
        )
    if "expected_clarification" in case and bool(
        payload.get("clarification_required")
    ) != case["expected_clarification"]:
        failures.append("clarification state mismatch")
    nodes = [row.get("node") for row in payload.get("pipeline_log", [])]
    for node in case.get("required_nodes", []):
        if node not in nodes:
            failures.append(f"missing agent node: {node}")
    if case.get("expect_no_agents") and nodes:
        failures.append(f"agents ran without required evidence: {nodes}")
    missing = payload.get("clarification", {}).get("missing_fields", [])
    for field in case.get("required_missing", []):
        if field not in missing:
            failures.append(f"missing clarification field: {field}")
    asset_results = [
        row.get("display_asset_id") for row in payload.get("multi_asset_results", [])
    ]
    if case.get("expected_asset_results") != asset_results and case.get(
        "expected_asset_results"
    ) is not None:
        failures.append(
            f"asset results {asset_results} != {case['expected_asset_results']}"
        )
    response_text = " ".join(
        [str(payload.get("response", ""))]
        + [str(item) for item in payload.get("details", [])]
        + [str(item) for item in payload.get("actions", [])]
    )
    if case.get("contains_any") and not any(
        item.lower() in response_text.lower() for item in case["contains_any"]
    ):
        failures.append(f"none of expected phrases found: {case['contains_any']}")
    return not failures, failures


def _write_workbook(rows: list[dict], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Chat QC"
    headers = [
        "ID", "Category", "Persona", "Question", "HTTP", "Intent",
        "Clarification", "Agent calls", "Evidence sources", "Response",
        "Actions", "HITL", "Reflection", "Latency ms", "Pass",
        "Failure reason", "Run ID", "Conversation ID", "Raw JSON",
    ]
    ws.append(headers)
    for row in rows:
        payload = row["payload"]
        nodes = [entry.get("node", "") for entry in payload.get("pipeline_log", [])]
        sources = [
            entry.get("data_sources", "") for entry in payload.get("pipeline_log", [])
            if entry.get("data_sources")
        ]
        hitl = payload.get("hitl") or {
            key: payload.get(key) for key in (
                "hitl_required", "hitl_remediation", "hitl_monitoring",
                "hitl_diagnosis", "hitl_knowledge", "hitl_advisory",
            ) if payload.get(key)
        }
        ws.append([
            row["case"]["id"], row["case"]["category"], row["case"]["persona"],
            row["case"]["question"], row["status_code"], payload.get("intent"),
            bool(payload.get("clarification_required")), ", ".join(nodes),
            "; ".join(sources), payload.get("response", ""),
            "\n".join(map(str, payload.get("actions", []))),
            json.dumps(hitl, ensure_ascii=False, default=str),
            f"{payload.get('reflection_status', '')} / {payload.get('reflection_iterations', '')}",
            row["latency_ms"], "PASS" if row["passed"] else "FAIL",
            "; ".join(row["failures"]), payload.get("run_id", ""),
            payload.get("conversation_id", ""),
            json.dumps(payload, ensure_ascii=False, default=str),
        ])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [12, 18, 16, 52, 8, 20, 14, 42, 48, 80, 55, 35, 22, 12, 10, 48, 24, 26, 90]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        cells[14].fill = PatternFill(
            "solid", fgColor="C6EFCE" if cells[14].value == "PASS" else "FFC7CE"
        )

    summary = wb.create_sheet("Summary")
    counts = Counter(row["case"]["category"] for row in rows)
    passed = sum(row["passed"] for row in rows)
    summary.append(["Metric", "Value"])
    summary.append(["Run timestamp UTC", datetime.now(timezone.utc).isoformat()])
    summary.append(["Cases", len(rows)])
    summary.append(["Passed", passed])
    summary.append(["Failed", len(rows) - passed])
    summary.append(["Pass rate", passed / len(rows) if rows else 0])
    summary.append(["Median calculation", "Use the Chat QC sheet Latency ms column"])
    for category, count in sorted(counts.items()):
        summary.append([f"Category: {category}", count])
    for cell in summary[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 42
    wb.save(output)


def _publish(rows: list[dict], dataset_name: str) -> str:
    from langsmith import Client

    client = Client()
    try:
        dataset = client.read_dataset(dataset_name=dataset_name)
    except Exception:
        dataset = client.create_dataset(
            dataset_name,
            description=(
                "DRO local /api/chat QC: hard routing, persona, evidence, HITL, "
                "multi-asset, clarification, memory, and fallback cases."
            ),
            metadata={"suite": "DRO_Chat_QC", "case_count": len(rows)},
        )
    examples = []
    for row in rows:
        case, payload = row["case"], row["payload"]
        examples.append({
            "inputs": {
                "question": case["question"], "persona": case["persona"],
                "context": case.get("context", {}),
            },
            "outputs": payload,
            "metadata": {
                "qc_id": case["id"], "category": case["category"],
                "passed_local_checks": row["passed"],
                "failures": row["failures"],
                "latency_ms": row["latency_ms"],
                "expected": {
                    key: value for key, value in case.items()
                    if key.startswith("expected_") or key.startswith("required_")
                },
            },
        })
    client.create_examples(dataset_id=dataset.id, examples=examples, max_concurrency=1)
    return str(dataset.id)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cases", type=Path, default=CASES_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--publish-langsmith", action="store_true")
    parser.add_argument("--dataset-name", default="DRO Chat QC 50")
    parser.add_argument("--disable-llm", action="store_true")
    args = parser.parse_args()
    # A local QC run must not wait for an unconfigured deployment database.
    # The API behavior is identical; HITL sessions use the repository's local
    # SQLite development adapter for this run.
    os.environ["HITL_REPOSITORY"] = "sqlite"
    if args.disable_llm:
        for key in ("AZURE_AI_KEY", "AZURE_OPENAI_KEY", "OPENAI_API_KEY"):
            os.environ[key] = ""
    sys.path.insert(0, str(ROOT))
    from fastapi.testclient import TestClient
    import src.api.main as api

    api.app.dependency_overrides[api.get_current_user] = lambda: {
        "sub": "local-qc", "display_name": "Local QC", "role": "admin",
        "allowed_personas": [
            "supervisor", "engineer", "maintenance", "manager",
            "executive", "md", "ot", "safety",
        ],
    }
    cases = json.loads(args.cases.read_text(encoding="utf-8"))
    if len(cases) != 50:
        raise ValueError(f"QC suite must contain exactly 50 cases, found {len(cases)}")
    rows = []
    with TestClient(api.app, raise_server_exceptions=False) as client:
        for index, case in enumerate(cases, 1):
            request = {
                "message": case["question"], "persona": case["persona"],
                "context": case.get("context"),
                "conversation_id": f"QC-{case['id']}-{int(time.time())}",
            }
            started = time.perf_counter()
            response = client.post("/api/chat", json=request)
            latency = round((time.perf_counter() - started) * 1000, 1)
            try:
                payload = response.json()
            except Exception:
                payload = {"response": response.text}
            passed, failures = _evaluate(case, response.status_code, payload)
            rows.append({
                "case": case, "status_code": response.status_code,
                "payload": payload, "latency_ms": latency,
                "passed": passed, "failures": failures,
            })
            print(
                f"[{index:02d}/50] {case['id']} "
                f"{'PASS' if passed else 'FAIL'} {latency:.0f}ms",
                flush=True,
            )
    _write_workbook(rows, args.output)
    passed = sum(row["passed"] for row in rows)
    print(f"Workbook: {args.output}")
    print(f"Result: {passed}/{len(rows)} passed local checks")
    if args.publish_langsmith:
        dataset_id = _publish(rows, args.dataset_name)
        print(f"LangSmith dataset: {args.dataset_name} ({dataset_id})")
    return 0 if passed == len(rows) else 1


if __name__ == "__main__":
    raise SystemExit(main())
