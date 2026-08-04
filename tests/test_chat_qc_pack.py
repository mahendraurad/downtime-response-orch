"""Certification for the repeatable 50-question chat QC artefacts."""
import json
from pathlib import Path

from openpyxl import load_workbook

from scripts.run_chat_qc import _evaluate, _write_workbook
from src.orchestrator.query_router import plan_query


ROOT = Path(__file__).resolve().parents[1]


def _cases():
    return json.loads(
        (ROOT / "quality" / "chat_qc_questions.json").read_text(encoding="utf-8")
    )


def test_qc_pack_has_50_unique_cases_and_all_personas():
    cases = _cases()
    assert len(cases) == 50
    assert len({case["id"] for case in cases}) == 50
    assert {case["persona"] for case in cases} == {
        "supervisor", "engineer", "maintenance", "manager",
        "executive", "ot", "safety",
    }


def test_qc_pack_covers_safety_critical_categories():
    categories = {case["category"] for case in _cases()}
    assert {
        "concept", "known_asset", "unknown_asset", "clarification",
        "fleet", "multi_asset", "persona", "memory", "adversarial",
        "evidence", "full_pipeline",
    } <= categories


def test_qc_evaluator_rejects_agent_calls_without_evidence():
    case = {"expected_clarification": True, "expect_no_agents": True}
    passed, failures = _evaluate(case, 200, {
        "clarification_required": True,
        "pipeline_log": [{"node": "predictive_risk"}],
    })
    assert not passed
    assert any("without required evidence" in failure for failure in failures)


def test_qc_workbook_contains_results_and_summary(tmp_path):
    case = {
        "id": "QC-T", "category": "concept", "persona": "engineer",
        "question": "What is RUL?",
    }
    output = tmp_path / "qc.xlsx"
    _write_workbook([{
        "case": case, "status_code": 200,
        "payload": {
            "intent": "concept", "response": "Remaining Useful Life",
            "pipeline_log": [], "clarification_required": False,
        },
        "latency_ms": 10.0, "passed": True, "failures": [],
    }], output)
    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames == ["Chat QC", "Summary"]
    assert workbook["Chat QC"]["A2"].value == "QC-T"
    assert workbook["Chat QC"]["O2"].value == "PASS"


def test_router_recognises_concepts_and_learning_phrasings():
    assert plan_query("Explain BPFO in plain English").intent == "concept"
    assert plan_query("How is BPFI different from BPFO?").intent == "concept"
    assert plan_query("Why do vibration and temperature matter?").intent == "concept"
    assert plan_query("Show the last three learned bearing cases").intent == "learning_history"


def test_router_recognises_complete_and_execution_phrasings():
    assert plan_query("Perform a complete technical analysis of M-104").pipeline_intent == "full"
    assert plan_query("Give an executive decision brief for M-104").pipeline_intent == "full"
    assert plan_query("Create a work order for Z-404").intent == "execution"
    assert plan_query("Can P-207 wait until next week?").intent == "risk"
